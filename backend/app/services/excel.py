from datetime import datetime
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import CellIsRule
from ..config import settings
COLORS={'SAFE':'C6EFCE','WARNING':'FFEB9C','ERROR':'FFC7CE','UNREACHABLE':'D9EAF7'}
def export(job, session, results):
    wb=Workbook(); summary=wb.active; summary.title='Summary'
    rows=[('Scan Time',job.start_time),('Cluster',session.cluster),('Teleport Proxy',session.proxy),('TSH Version',session.tsh_version),('Total Nodes',job.total_nodes),('Safe',job.safe_nodes),('Warning',job.warning_nodes),('Error',job.error_nodes),('Unreachable',job.unreachable_nodes)]
    for r in rows: summary.append(r)
    ws=wb.create_sheet('Scan Results'); headers=['STT','Status','Hostname','IP','PID','PPID','User','Executable','Command','Runtime','CGroup','Risk','Scan Time']; ws.append(headers)
    er=wb.create_sheet('Errors'); er.append(['Hostname','IP','Error Type','Error Message','Scan Time'])
    for i,x in enumerate(results,1):
        ws.append([i,x.status,x.hostname,x.address,x.pid,x.ppid,x.username,x.executable,x.command_line,x.runtime,x.cgroup_path,x.risk,x.scanned_at])
        if x.status in ('ERROR','UNREACHABLE'): er.append([x.hostname,x.address,x.status,x.error_message,x.scanned_at])
    for sheet in (summary,ws,er):
        sheet.freeze_panes='A2'; sheet.auto_filter.ref=sheet.dimensions
        for c in sheet[1]: c.font=Font(bold=True); c.fill=PatternFill('solid',fgColor='1F4E78'); c.font=Font(bold=True,color='FFFFFF')
        for col in sheet.columns: sheet.column_dimensions[col[0].column_letter].width=min(60,max(12,max(len(str(v.value or '')) for v in col)+2));
        for row in sheet.iter_rows():
            for c in row: c.alignment=Alignment(vertical='top',wrap_text=True)
    for status,color in COLORS.items(): ws.conditional_formatting.add(f'B2:B1048576',CellIsRule(operator='equal',formula=[f'"{status}"'],fill=PatternFill('solid',fgColor=color)))
    name=f'teleport_cgroup_scan_{datetime.utcnow():%Y%m%d_%H%M%S}.xlsx'; p=settings.export_root/name; wb.save(p); return p
