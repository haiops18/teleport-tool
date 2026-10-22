from datetime import datetime, timedelta
import re
from pathlib import Path
from sqlalchemy import select, func, case, and_, or_
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from ..models import Node, ScanJob, ScanResult, ServiceConfigResult
from ..config import settings

def status_of(value):
    return value if value in {'SAFE','WARNING','ERROR'} else 'UNKNOWN'
def service_config_of(db,node_id):
    value=db.scalar(select(ServiceConfigResult).where(ServiceConfigResult.node_id==node_id).order_by(ServiceConfigResult.id.desc()))
    return {'kill_mode':(value.new_kill_mode or value.previous_kill_mode) if value else None,'service_config_status':value.status if value else 'UNKNOWN','service_config_checked_at':value.finished_at if value else None,'restart_status':value.restart_status if value else None}
def time_filter(value):
    now=datetime.utcnow(); options={'today':now.replace(hour=0,minute=0,second=0,microsecond=0),'24h':now-timedelta(hours=24),'7d':now-timedelta(days=7),'30d':now-timedelta(days=30)}
    return options.get(value)
def scoped_results(session_id, since=None):
    q=select(ScanResult).join(ScanJob,ScanResult.scan_job_id==ScanJob.id).where(ScanJob.session_id==session_id)
    return q.where(ScanResult.scanned_at>=since) if since else q
def summary(db, session_id, period=None):
    since=time_filter(period); rows=db.scalars(scoped_results(session_id,since)).all(); hosts={x.node_id for x in rows if x.node_id}; latest={}
    for x in rows:
        if x.node_id and (x.node_id not in latest or x.scanned_at>latest[x.node_id].scanned_at): latest[x.node_id]=x
    statuses=[status_of(x.status) for x in latest.values()]
    jobs=db.scalar(select(func.count()).select_from(ScanJob).where(ScanJob.session_id==session_id)) or 0
    return {'total_hosts_scanned':len(hosts),'safe_hosts':statuses.count('SAFE'),'warning_hosts':statuses.count('WARNING'),'error_hosts':statuses.count('ERROR'),'total_scan_jobs':jobs,'successful_scans':sum(x.status=='SAFE' for x in rows),'warning_scans':sum(x.status=='WARNING' for x in rows),'failed_scans':sum(x.status in {'ERROR','UNREACHABLE'} for x in rows),'total_processes_found':sum(x.pid is not None for x in rows),'risk_processes':sum(x.classification=='RISK' for x in rows),'protected_external_processes':sum(x.classification=='EXTERNAL_PROTECTED' for x in rows),'last_scan':max((x.scanned_at for x in rows),default=None)}
def hosts(db, session_id, page=1, page_size=50, status=None, search=None, period=None):
    q=select(Node).where(Node.session_id==session_id)
    if search: q=q.where(or_(Node.hostname.contains(search),Node.address.contains(search),Node.node_id.contains(search)))
    nodes=db.scalars(q.order_by(Node.hostname)).all(); since=time_filter(period); items=[]
    for n in nodes:
        results=db.scalars(scoped_results(session_id,since).where(ScanResult.node_id==n.id).order_by(ScanResult.scanned_at.desc())).all()
        if not results: continue
        latest=results[0]; current=status_of(latest.status)
        if status and current!=status: continue
        scans=len({x.scan_job_id for x in results}); processes=sum(x.pid is not None for x in results); risks=sum(x.classification=='RISK' for x in results); protected=sum(x.classification=='EXTERNAL_PROTECTED' for x in results)
        items.append({'node_id':n.id,'hostname':n.hostname,'node_name':n.node_id or n.hostname,'address':n.address,'labels':n.labels or {},'status':current,'last_scan':latest.scanned_at,'total_scans':scans,'processes_found':processes,'risk_processes':risks,'protected_external_processes':protected,'kill_mode':latest.kill_mode or service_config_of(db,n.id)['kill_mode'],'last_result':latest.error_message or current,**service_config_of(db,n.id)})
    total=len(items); start=(page-1)*page_size
    return {'items':items[start:start+page_size],'page':page,'page_size':page_size,'total':total}
def host_detail(db, session_id, node_id):
    node=db.scalar(select(Node).where(Node.id==node_id,Node.session_id==session_id))
    if not node: return None
    rows=db.scalars(scoped_results(session_id).where(ScanResult.node_id==node_id).order_by(ScanResult.scanned_at.desc())).all()
    if not rows: return {'node_id':node.id,'hostname':node.hostname,'address':node.address,'labels':node.labels or {},'status':'UNKNOWN','total_scans':0,'history':[],**service_config_of(db,node.id)}
    by_job={}
    for r in rows: by_job.setdefault(r.scan_job_id,[]).append(r)
    history=[]
    for jid, rs in by_job.items():
        job=db.get(ScanJob,jid); last=max(rs,key=lambda x:x.scanned_at); history.append({'scan_id':jid,'scan_time':last.scanned_at,'status':status_of(last.status),'processes':sum(x.pid is not None for x in rs),'risk_processes':sum(x.classification=='RISK' for x in rs),'protected_processes':sum(x.classification=='EXTERNAL_PROTECTED' for x in rs),'kill_mode':last.kill_mode,'user':job.created_by if job else None,'error':last.error_message,'duration_seconds':((job.end_time-job.start_time).total_seconds() if job and job.end_time else None)})
    history.sort(key=lambda x:x['scan_time'],reverse=True)
    return {'node_id':node.id,'hostname':node.hostname,'node_name':node.node_id or node.hostname,'address':node.address,'labels':node.labels or {},'first_scan':history[-1]['scan_time'],'last_scan':history[0]['scan_time'],'status':history[0]['status'],'total_scans':len(history),'successful_scans':sum(x['status']=='SAFE' for x in history),'warning_scans':sum(x['status']=='WARNING' for x in history),'failed_scans':sum(x['status']=='ERROR' for x in history),'processes_found':sum(x['processes'] for x in history),'risk_processes':sum(x['risk_processes'] for x in history),'history':history,**service_config_of(db,node.id)}
def scan_detail(db, session_id, scan_id):
    job=db.scalar(select(ScanJob).where(ScanJob.id==scan_id,ScanJob.session_id==session_id))
    if not job: return None
    rows=db.scalars(select(ScanResult).where(ScanResult.scan_job_id==scan_id).order_by(ScanResult.hostname,ScanResult.pid)).all()
    return {'scan_id':job.id,'started_at':job.start_time,'finished_at':job.end_time,'duration_seconds':((job.end_time-job.start_time).total_seconds() if job.end_time else None),'status':job.status,'teleport_user':job.created_by,'kill_mode':next((x.kill_mode for x in rows if x.kill_mode),None),'processes':[{'risk':'RISK' if x.classification=='RISK' else ('PROTECTED' if x.classification=='EXTERNAL_PROTECTED' else 'SAFE'),'classification':x.classification or 'UNKNOWN','pid':x.pid,'ppid':x.ppid,'user':x.username,'process_name':x.executable,'command_line':x.command_line,'cgroup':x.cgroup_path,'detected_at':x.scanned_at,'reason':x.risk or x.error_message,'hostname':x.hostname} for x in rows]}
def safe_excel(v):
    if not isinstance(v,str): return v
    v=re.sub(r'[\x00-\x08\x0B\x0C\x0E-\x1F]','',v)
    return "'"+v if v[:1] in '=+-@' else v
def export_report(db, session_id, node_id=None, node_ids=None, search=None, status=None):
    report=host_detail(db,session_id,node_id) if node_id else None; rows=(report['history'] if report else hosts(db,session_id,page_size=500,status=status,search=search)['items']);
    if node_ids is not None: rows=[row for row in rows if row['node_id'] in set(node_ids)]
    wb=Workbook(); ws=wb.active; ws.title='01_Summary';
    for k,v in summary(db,session_id).items(): ws.append([k.replace('_',' ').title(),v])
    hs=wb.create_sheet('02_Hosts'); hs.append(['No','Hostname','Node Name','IP','KillMode','Status','Last Scan','Total Scans','Processes','Risk Processes','Protected External Processes'])
    host_rows=[report] if report else rows
    for i,x in enumerate(host_rows,1): hs.append([i,safe_excel(x['hostname']),safe_excel(x.get('node_name','')),safe_excel(x.get('address','')),x.get('kill_mode'),x.get('status'),x.get('last_scan'),x.get('total_scans'),x.get('processes_found'),x.get('risk_processes'),x.get('protected_external_processes',0)])
    sh=wb.create_sheet('03_Scan_History'); sh.append(['No','Scan ID','Hostname','Scan Time','KillMode','Status','Duration','Process Count','Risk Process Count','Protected Process Count','Error'])
    processes=wb.create_sheet('04_Process_Details'); processes.append(['No','Scan ID','Hostname','PID','PPID','User','Process','Command','Cgroup','Classification','Risk','Reason'])
    index=1
    for host in host_rows:
        detail=host_detail(db,session_id,host['node_id'])
        for h in detail['history']:
            sh.append([index,h['scan_id'],safe_excel(detail['hostname']),h['scan_time'],h.get('kill_mode'),h['status'],h['duration_seconds'],h['processes'],h['risk_processes'],h.get('protected_processes',0),safe_excel(h['error'] or '')]); scan=scan_detail(db,session_id,h['scan_id'])
            for p in scan['processes']: processes.append([index,h['scan_id'],safe_excel(p['hostname']),p['pid'],p['ppid'],safe_excel(p['user'] or ''),safe_excel(p['process_name'] or ''),safe_excel(p['command_line'] or ''),safe_excel(p['cgroup'] or ''),p.get('classification'),p['risk'],safe_excel(p['reason'] or '')])
            index+=1
    for sheet in wb.worksheets:
        sheet.freeze_panes='A2'; sheet.auto_filter.ref=sheet.dimensions
        for c in sheet[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='1F4E78')
        for col in sheet.columns: sheet.column_dimensions[col[0].column_letter].width=min(60,max(12,max(len(str(c.value or '')) for c in col)+2))
        for row in sheet.iter_rows():
            for c in row: c.alignment=Alignment(vertical='top',wrap_text=True)
    name=f'Teleport_Cgroup_Report_{datetime.utcnow():%Y%m%d_%H%M%S}.xlsx'; path=Path(settings.export_root)/name; wb.save(path); return path
