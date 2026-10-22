"""Read-only queries and Excel export for Compute Discovery history."""
from collections import defaultdict
from datetime import datetime
from pathlib import Path
import re
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.chart import BarChart, PieChart, Reference
from sqlalchemy import String, cast, func, or_, select
from ..config import settings
from ..models import ComputeDiscoveryJob, ComputeDiscoveryResult, Node


def _clean_cell(value):
    text='' if value is None else str(value)
    # Excel rejects control bytes (for example ANSI escape sequences emitted by tsh).
    text=re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', text)
    return "'"+text if text.startswith(('=','+','-','@')) else text


def _latest_stmt(session_id):
    latest_ids=(select(func.max(ComputeDiscoveryResult.id).label('id'))
        .where(ComputeDiscoveryResult.session_id==session_id)
        .group_by(ComputeDiscoveryResult.node_id).subquery())
    return (select(ComputeDiscoveryResult, ComputeDiscoveryJob, Node)
        .join(ComputeDiscoveryJob,ComputeDiscoveryJob.id==ComputeDiscoveryResult.job_id)
        .outerjoin(Node,Node.id==ComputeDiscoveryResult.node_id)
        .where(ComputeDiscoveryResult.id.in_(select(latest_ids.c.id))))


def _filtered(stmt, search=None, status=None, job_id=None):
    if job_id: stmt=stmt.where(ComputeDiscoveryResult.job_id==job_id)
    if status: stmt=stmt.where(ComputeDiscoveryResult.status==status.upper())
    if search:
        value=f'%{search.strip()}%'
        stmt=stmt.where(or_(ComputeDiscoveryResult.compute.ilike(value),ComputeDiscoveryResult.hostname.ilike(value),ComputeDiscoveryResult.address.ilike(value),ComputeDiscoveryJob.job_name.ilike(value),cast(Node.labels,String).ilike(value)))
    return stmt


def _row(result, job, node):
    return {'node_id':result.node_id,'hostname':result.hostname,'address':result.address,'labels':node.labels if node else None,
        'nic':result.nic,'compute':result.compute,'status':result.status,'failed_stage':result.failed_stage,'command_exit_code':result.command_exit_code,'reason_code':result.reason_code or ('LEGACY_UNKNOWN' if result.status=='NOT_DETECTED' else None),'reason_message':result.reason_message or ('No detailed reason was stored for this historical result.' if result.status=='NOT_DETECTED' else None),'detail':result.detail,'discovered_at':result.discovered_at,
        'job_id':result.job_id,'job_name':job.job_name,'job_status':job.status}


def _latest_rows(db, session_id, search=None, status=None, job_id=None):
    return [_row(*item) for item in db.execute(_filtered(_latest_stmt(session_id),search,status,job_id).order_by(ComputeDiscoveryResult.hostname)).all()]


def _page(items, page, page_size):
    total=len(items); start=(page-1)*page_size
    return {'items':items[start:start+page_size],'total':total,'page':page,'page_size':page_size}


def summary(db, session_id):
    rows=_latest_rows(db,session_id)
    detected=sum(x['status'] in ('DETECTED','SUCCESS') for x in rows); not_detected=sum(x['status'] in ('NOT_DETECTED','NO_LLDP') for x in rows)
    distribution={'DETECTED':detected,'NOT_DETECTED':not_detected,'UNREACHABLE':sum(x['status']=='UNREACHABLE' for x in rows),'ERROR':sum(x['status']=='ERROR' for x in rows)}
    reasons=defaultdict(int)
    for x in rows:
        if x['status'] in ('NOT_DETECTED','NO_LLDP'): reasons[x['reason_code'] or 'LEGACY_UNKNOWN']+=1
    return {'total_nodes_scanned':len(rows),'total_nodes':len(rows),'detected':detected,'not_detected':not_detected,'unreachable':distribution['UNREACHABLE'],'error':distribution['ERROR'],'compute_found':detected,
        'unique_computes':len({x['compute'] for x in rows if x['compute']}),'detection_rate':round(detected*100/len(rows),1) if rows else 0,
        'status_distribution':distribution,'not_detected_reason_distribution':[{'reason_code':k,'count':v} for k,v in sorted(reasons.items(),key=lambda pair:(-pair[1],pair[0]))],
        'items':rows}


def groups(db, session_id, page, page_size, search=None, status=None, job_id=None):
    bucket=defaultdict(list)
    for row in _latest_rows(db,session_id,search,status,job_id):
        key=row['compute'] if row['status'] in ('DETECTED','SUCCESS') and row['compute'] else ('NOT_DETECTED' if row['status'] in ('NOT_DETECTED','NO_LLDP') else row['status'])
        bucket[key].append(row)
    result=[]
    for compute, members in bucket.items():
        latest=max((x['discovered_at'] for x in members if x['discovered_at']),default=None)
        result.append({'compute':compute,'node_count':len(members),'nodes':members,'node_names':[x['hostname'] for x in members],
            'node_ips':[x['address'] for x in members if x['address']], 'status':'DETECTED' if compute not in ('NOT_DETECTED','ERROR','UNREACHABLE') else compute,
            'last_scan':latest,'job_name':members[0]['job_name']})
    result.sort(key=lambda x:(-x['node_count'],x['compute'].lower()))
    return _page(result,page,page_size)


def nodes(db, session_id, page, page_size, search=None, status=None, job_id=None):
    return _page(_latest_rows(db,session_id,search,status,job_id),page,page_size)


def history(db, session_id, page, page_size, search=None, status=None, job_id=None):
    stmt=(select(ComputeDiscoveryResult,ComputeDiscoveryJob,Node)
        .join(ComputeDiscoveryJob,ComputeDiscoveryJob.id==ComputeDiscoveryResult.job_id)
        .outerjoin(Node,Node.id==ComputeDiscoveryResult.node_id)
        .where(ComputeDiscoveryResult.session_id==session_id))
    rows=[_row(*item) for item in db.execute(_filtered(stmt,search,status,job_id).order_by(ComputeDiscoveryResult.discovered_at.desc(),ComputeDiscoveryResult.id.desc())).all()]
    return _page(rows,page,page_size)


def job_results(db, session_id, job_id):
    stmt=(select(ComputeDiscoveryResult,ComputeDiscoveryJob,Node)
        .join(ComputeDiscoveryJob,ComputeDiscoveryJob.id==ComputeDiscoveryResult.job_id)
        .outerjoin(Node,Node.id==ComputeDiscoveryResult.node_id)
        .where(ComputeDiscoveryResult.session_id==session_id,ComputeDiscoveryResult.job_id==job_id)
        .order_by(ComputeDiscoveryResult.discovered_at.desc()))
    return [_row(*item) for item in db.execute(stmt).all()]


def _sheet(book, name, headers, rows):
    ws=book.create_sheet(name); ws.append(headers)
    for cell in ws[1]: cell.font=Font(bold=True)
    for row in rows: ws.append([_clean_cell(value) for value in row])
    for column in ws.columns:
        letter=column[0].column_letter; ws.column_dimensions[letter].width=min(60,max(12,max(len(str(c.value or '')) for c in column)+2))
        for cell in column: cell.alignment=Alignment(wrap_text=True,vertical='top')
    return ws


def export_discovery(db, session, search=None, status=None, job_id=None, node_ids=None):
    latest=_latest_rows(db,session.id,search,status,job_id)
    if node_ids is not None: latest=[x for x in latest if x['node_id'] in set(node_ids)]
    group_rows=groups(db,session.id,1,100000,search,status,job_id)['items']
    history_rows=history(db,session.id,1,100000,search,status,job_id)['items']
    if node_ids is not None: history_rows=[x for x in history_rows if x['node_id'] in set(node_ids)]
    jobs=db.scalars(select(ComputeDiscoveryJob).where(ComputeDiscoveryJob.session_id==session.id).order_by(ComputeDiscoveryJob.created_at.desc())).all()
    book=Workbook(); book.remove(book.active)
    overview=_sheet(book,'01_Overview',['Metric','Value'],[
        ('Generated At',datetime.utcnow().isoformat()),('Generated By',session.username),('Session / Cluster',session.cluster or ''),('Filter Summary',f'search={search or ""}; status={status or ""}; job={job_id or ""}'),
        ('Total Nodes',len(latest)),('Detected',sum(x['status'] in ('DETECTED','SUCCESS') for x in latest)),('Not Detected',sum(x['status'] in ('NOT_DETECTED','NO_LLDP') for x in latest)),('Unreachable',sum(x['status']=='UNREACHABLE' for x in latest)),('Error',sum(x['status']=='ERROR' for x in latest)),('Unique Computes',len({x['compute'] for x in latest if x['compute']})),('Detection Rate',f"{round(sum(x['status'] in ('DETECTED','SUCCESS') for x in latest)*100/len(latest),1) if latest else 0}%")])
    groups_sheet=_sheet(book,'02_Compute_Groups',['Compute Name','Node Count','Node Names','Node IPs','Systems','Latest Discovery Job','Latest Scan','Status'],[(x['compute'],x['node_count'],'\n'.join(x['node_names']),'\n'.join(x['node_ips']),'',x['job_name'],x['last_scan'],x['status']) for x in group_rows if x['status']=='DETECTED'])
    _sheet(book,'03_Nodes',['Node Name','IP','System','Labels','Compute Name','Discovery Status','Reason Code','Reason','Job Name','Job ID','Scanned At'],[(x['hostname'],x['address'],'',x['labels'],x['compute'],x['status'],x['reason_code'],x['reason_message'],x['job_name'],x['job_id'],x['discovered_at']) for x in latest])
    _sheet(book,'04_Not_Detected',['Node Name','IP','System','Labels','Reason Code','Reason','Job Name','Job ID','Scanned At','Last Known Compute'],[(x['hostname'],x['address'],'',x['labels'],x['reason_code'],x['reason_message'],x['job_name'],x['job_id'],x['discovered_at'],'') for x in latest if x['status'] in ('NOT_DETECTED','NO_LLDP')])
    _sheet(book,'05_Discovery_History',['Job ID','Job Name','Node Name','IP','Compute','Status','Reason Code','Reason','Raw LLDP Summary','Scanned At'],[(x['job_id'],x['job_name'],x['hostname'],x['address'],x['compute'],x['status'],x['reason_code'],x['reason_message'],x['detail'],x['discovered_at']) for x in history_rows])
    _sheet(book,'06_Job_Details',['Job ID','Job Name','Status','Created By','Created At','Started At','Finished At','Total Nodes','Completed','Detected','Not Detected','Unreachable','Error'],[(x.id,x.job_name,x.status,x.created_by,x.created_at,x.started_at,x.finished_at,x.total_nodes,x.completed_nodes,x.success_nodes,'','',x.failed_nodes) for x in jobs])
    status_start=overview.max_row+2; overview.append(['Status','Count'])
    for label in ('DETECTED','NOT_DETECTED','UNREACHABLE','ERROR'): overview.append([label,sum(x['status']==label or (label=='DETECTED' and x['status']=='SUCCESS') or (label=='NOT_DETECTED' and x['status']=='NO_LLDP') for x in latest)])
    pie=PieChart(); pie.title='Discovery Status Distribution'; pie.add_data(Reference(overview,min_col=2,min_row=status_start,max_row=status_start+4),titles_from_data=True); pie.set_categories(Reference(overview,min_col=1,min_row=status_start+1,max_row=status_start+4)); overview.add_chart(pie,'D2')
    compute_start=overview.max_row+2; overview.append(['Compute','Node Count'])
    for row in group_rows[:10]: overview.append([row['compute'],row['node_count']])
    bar=BarChart(); bar.type='bar'; bar.title='Top Compute by Node Count'; bar.add_data(Reference(overview,min_col=2,min_row=compute_start,max_row=overview.max_row),titles_from_data=True); bar.set_categories(Reference(overview,min_col=1,min_row=compute_start+1,max_row=overview.max_row)); overview.add_chart(bar,'D18')
    reason_start=overview.max_row+2; overview.append(['Reason Code','Count'])
    for code,count in sorted(defaultdict(int, {x['reason_code'] or 'LEGACY_UNKNOWN':sum(y['status'] in ('NOT_DETECTED','NO_LLDP') and (y['reason_code'] or 'LEGACY_UNKNOWN')==(x['reason_code'] or 'LEGACY_UNKNOWN') for y in latest) for x in latest if x['status'] in ('NOT_DETECTED','NO_LLDP')}).items()): overview.append([code,count])
    reason_chart=BarChart(); reason_chart.type='bar'; reason_chart.title='NOT_DETECTED Reasons'; reason_chart.add_data(Reference(overview,min_col=2,min_row=reason_start,max_row=overview.max_row),titles_from_data=True); reason_chart.set_categories(Reference(overview,min_col=1,min_row=reason_start+1,max_row=overview.max_row)); overview.add_chart(reason_chart,'D34')
    settings.export_root.mkdir(parents=True,exist_ok=True); output=settings.export_root/f'compute_discovery_{datetime.utcnow().strftime("%Y%m%d_%H%M%S")}.xlsx'; book.save(output)
    check=load_workbook(output)
    expected=['01_Overview','02_Compute_Groups','03_Nodes','04_Not_Detected','05_Discovery_History','06_Job_Details']
    if check.sheetnames!=expected or len(check['01_Overview']._charts)!=3 or any(tuple(next(check[name].iter_rows(values_only=True)))==() for name in expected): raise RuntimeError('DISCOVERY_EXPORT_VALIDATION_FAILED')
    return output
